import pandas as pd
import os
import re
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

def normalize_column_name(column_name):
    """
    统一表头名称，兼容“1、姓名”这类问卷星导出格式。
    """
    normalized = str(column_name).strip()
    normalized = re.sub(r'^\d+\s*[、.．)\]】）]?\s*', '', normalized)
    return normalized.strip()

def normalize_cell_text(value):
    """
    清理普通文本字段，避免空值被转成 "nan" 参与去重。
    """
    if pd.isna(value):
        return ''
    return str(value).replace('\u3000', ' ').strip()

def normalize_student_id(value):
    """
    统一学号格式。

    Excel 里学号有时会被 pandas 读成 25215150414.0 或科学计数法，
    这里统一还原成纯数字文本，确保两个表能正确匹配。
    """
    text = normalize_cell_text(value).replace(' ', '')
    if text.lower() in {'nan', 'none', 'null'}:
        return ''
    if re.fullmatch(r'\d+\.0+', text):
        return text.split('.')[0]
    if re.fullmatch(r'\d+(?:\.\d+)?[eE][+-]?\d+', text):
        try:
            return format(float(text), '.0f')
        except ValueError:
            return text
    return re.sub(r'\.0+$', '', text)

def save_dataframe_with_template(df, output_file, template_file, description):
    """
    使用模板文件导出 DataFrame，尽量保留原始格式。
    """
    print(f"[导出] 正在导出{description} ({output_file})...")
    try:
        wb = load_workbook(template_file)
        ws = wb.active

        # 清空原有数据（保留表头）
        if ws.max_row > 1:
            ws.delete_rows(2, amount=ws.max_row - 1)

        # 从第 2 行开始写入数据，保留模板中的表头
        rows = dataframe_to_rows(df, index=False, header=False)
        for r_idx, row in enumerate(rows, start=2):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)

        wb.save(output_file)
        print(f"[完成] {description}已保存至: {output_file}")

    except Exception as e:
        print(f"[警告] {description}格式保存失败（{e}），尝试直接保存为普通Excel...")
        df.to_excel(output_file, index=False)
        print(f"[完成] {description}已保存(无格式版本): {output_file}")

def merge_and_deduplicate(input_files, output_file, subset_cols=['姓名', '学号'], template_file=None, duplicate_output_file=None):
    """
    合并多个Excel文件并去重，保留第一个文件的格式。
    
    :param input_files: 文件路径列表
    :param output_file: 输出文件名
    :param subset_cols: 去重依据的列名列表
    :param template_file: 模板文件（用于保留格式），默认使用列表中的第一个文件
    :param duplicate_output_file: 重复名单输出文件名，默认不单独导出
    """
    print(f"--- 开始处理表格 (共 {len(input_files)} 个文件) ---")
    
    all_dfs = []
    valid_files = []
    
    # 1. 循环读取所有文件
    for file in input_files:
        if not os.path.exists(file):
            print(f"[警告] 找不到文件 {file}，已跳过")
            continue
            
        print(f"[读取] 正在读取: {file}...")
        try:
            df = pd.read_excel(file, dtype=str)
            # 数据预处理：统一表头、清理去重字段中的空格，防止匹配失败
            df.columns = [normalize_column_name(c) for c in df.columns]
            for col in subset_cols:
                if col in df.columns:
                    if col == '学号':
                        df[col] = df[col].apply(normalize_student_id)
                    else:
                        df[col] = df[col].apply(normalize_cell_text)

            # 去掉姓名/学号为空的空白行，否则空行之间也会被当作重复人员。
            present_subset_cols = [col for col in subset_cols if col in df.columns]
            if present_subset_cols:
                before_drop_blank = len(df)
                df = df[~df[present_subset_cols].eq('').any(axis=1)].copy()
                dropped_blank = before_drop_blank - len(df)
                if dropped_blank:
                    print(f"   - 已跳过 {dropped_blank} 行姓名或学号为空的数据")
            
            all_dfs.append(df)
            valid_files.append(file)
        except Exception as e:
            print(f"[错误] 读取 {file} 失败: {e}")

    if not all_dfs:
        print("[错误] 没有有效的数据可以合并")
        return

    # 2. 合并数据
    print("[合并] 正在合并数据...")
    combined_df = pd.concat(all_dfs, ignore_index=True)
    original_count = len(combined_df)

    # 3. 执行去重
    # 检查去重列是否存在
    missing_cols = [c for c in subset_cols if c not in combined_df.columns]
    if missing_cols:
        print(f"[错误] 表格中缺少去重必要的列 {missing_cols}")
        print(f"[提示] 当前表头包含: {list(combined_df.columns)}")
        return

    print(f"[去重] 正在根据 {subset_cols} 进行去重...")
    duplicate_people_df = combined_df[combined_df.duplicated(subset=subset_cols, keep=False)].copy()
    duplicate_people_df = duplicate_people_df.drop_duplicates(subset=subset_cols, keep='first')
    clean_df = combined_df.drop_duplicates(subset=subset_cols, keep='first')
    
    final_count = len(clean_df)
    print(f"[统计] 统计结果:")
    print(f"   - 原始总行数: {original_count}")
    print(f"   - 去重后行数: {final_count}")
    print(f"   - 剔除重复项: {original_count - final_count} 条")
    print(f"   - 重复人员数: {len(duplicate_people_df)}")

    # 4. 写入文件并保留格式
    # 默认使用第一个有效文件作为样式模板
    if not template_file:
        template_file = valid_files[0]

    save_dataframe_with_template(clean_df, output_file, template_file, "去重后的名单")

    if duplicate_output_file:
        if duplicate_people_df.empty:
            print("[提示] 未发现重复人员，仍会生成一个仅含表头的重复名单文件。")
        save_dataframe_with_template(duplicate_people_df, duplicate_output_file, template_file, "重复名单")

# ==========================================
# 🔧 配置区域
# ==========================================
if __name__ == "__main__":
    # 1. 待合并的文件列表 (请确保这些文件在当前目录下)
    FILES_TO_MERGE = [
        '第一次问卷星(1).xlsx',
        '第一次参与人员加分表.xlsx',
        # '可以继续添加更多文件名.xlsx'
    ]
    
    # 2. 输出文件名
    OUTPUT_NAME = '合并去重后的最终名单.xlsx'
    DUPLICATE_OUTPUT_NAME = '重复名单.xlsx'
    
    # 3. 判定重复的依据（这两列都一样才算重复）
    DEDUP_COLUMNS = ['姓名', '学号']

    merge_and_deduplicate(
        FILES_TO_MERGE,
        OUTPUT_NAME,
        DEDUP_COLUMNS,
        duplicate_output_file=DUPLICATE_OUTPUT_NAME
    )
